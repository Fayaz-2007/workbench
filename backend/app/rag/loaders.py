"""Document loaders — turn a file on disk into a `LoadedDocument`.

One small loader per format, dispatched by extension in `load_document()`.
A PDF page with almost no extractable text is flagged `needs_ocr=True`
rather than guessed at — `RAGService` (see `service.py`) is what decides
whether to run it through the OCR pipeline (`app/multimodal/`).
"""

from __future__ import annotations

import csv
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

from app.core.exceptions import InvalidRequestError
from app.rag.schemas import DocumentPage, LoadedDocument

# Below this many characters, a PDF page is treated as "no extractable
# text" (i.e. almost certainly a scan) and flagged for OCR instead of
# being ingested as an (empty) text page.
_MIN_TEXT_CHARS_PER_PAGE = 20
_MAX_ROWS_PREVIEWED = 200


def _new_document(filename: str, source_path: str, doc_type: str, pages: list[DocumentPage]) -> LoadedDocument:
    return LoadedDocument(
        document_id=uuid4().hex,
        filename=filename,
        source_path=source_path,
        doc_type=doc_type,
        pages=pages,
        ingested_at=datetime.now(UTC).isoformat(),
    )


def load_txt(path: Path) -> LoadedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    return _new_document(path.name, str(path), "txt", [DocumentPage(text=text)])


def load_pdf(path: Path) -> LoadedDocument:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[DocumentPage] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        pages.append(DocumentPage(text=text, page_number=i, needs_ocr=len(text) < _MIN_TEXT_CHARS_PER_PAGE))
    return _new_document(path.name, str(path), "pdf", pages)


def load_docx(path: Path) -> LoadedDocument:
    import docx

    document = docx.Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
    return _new_document(path.name, str(path), "docx", [DocumentPage(text=text)])


def load_csv(path: Path) -> LoadedDocument:
    with path.open(encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.reader(f)
        rows = [row for _, row in zip(range(_MAX_ROWS_PREVIEWED), reader)]
    text = "\n".join(", ".join(cell for cell in row) for row in rows)
    return _new_document(path.name, str(path), "csv", [DocumentPage(text=text)])


def load_image(path: Path) -> LoadedDocument:
    """Runs the configured OCR provider (see `app/multimodal/ocr.py`) over
    a standalone image file. If OCR is unavailable (e.g. the `tesseract`
    binary isn't installed), this still returns a `LoadedDocument` — with
    empty text and the reason visible in server logs — rather than
    failing ingestion outright.

    Note: this does not rasterize scanned *PDF* pages to images (that
    needs a PDF-to-image dependency such as PyMuPDF/pdf2image, not
    included here) — `load_pdf` instead flags such pages `needs_ocr=True`
    and leaves their text empty. Standalone image uploads (PNG/JPG) are
    fully supported.
    """
    from app.multimodal.ocr import get_ocr_provider

    result = get_ocr_provider().extract(path)
    page = DocumentPage(text=result.text, needs_ocr=not result.ok)
    return _new_document(path.name, str(path), "image", [page])


def load_xlsx(path: Path) -> LoadedDocument:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    pages: list[DocumentPage] = []
    for sheet in workbook.worksheets:
        lines = [f"# Sheet: {sheet.title}"]
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= _MAX_ROWS_PREVIEWED:
                break
            lines.append(", ".join("" if cell is None else str(cell) for cell in row))
        pages.append(DocumentPage(text="\n".join(lines)))
    workbook.close()
    return _new_document(path.name, str(path), "xlsx", pages)


_LOADERS = {
    ".txt": load_txt,
    ".md": load_txt,
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".csv": load_csv,
    ".xlsx": load_xlsx,
    ".png": load_image,
    ".jpg": load_image,
    ".jpeg": load_image,
}


def supported_extensions() -> list[str]:
    return sorted(_LOADERS)


def load_document(path: Path, display_filename: str | None = None) -> LoadedDocument:
    """`display_filename` overrides the on-disk name in the returned
    `LoadedDocument.filename` — used when `path` is an internal, uuid
    -prefixed storage name (see `app/api/routes/files.py`) but citations
    and document listings should show the user's original filename.
    """
    loader = _LOADERS.get(path.suffix.lower())
    if loader is None:
        raise InvalidRequestError(
            f"Unsupported document type '{path.suffix}'. Supported: {', '.join(supported_extensions())}.",
            extension=path.suffix,
        )
    document = loader(path)
    if display_filename:
        document.filename = display_filename
    return document


def load_bytes(filename: str, data: bytes, tmp_dir: Path) -> LoadedDocument:
    """Convenience for callers holding raw bytes (an upload) rather than an
    on-disk path: writes to `tmp_dir` under `filename` first, since every
    loader above works from a real file (pypdf/python-docx/openpyxl all
    expect a path or file-like — writing once keeps this simple and lets
    the same bytes double as the persisted copy in `documents_dir`).
    """
    target = tmp_dir / filename
    target.write_bytes(data)
    return load_document(target)


__all__ = ["load_document", "load_bytes", "supported_extensions"]
