"""Natural document/report queries must route to the Document Agent and
actually retrieve from the Knowledge Database — regression coverage for
the bug where "say about the report" (a single, unambiguous "report" hit)
fell through the Task Router's confidence floor to the General Agent,
which has no RAG capability at all and produced an ungrounded mock reply.

Also covers: DOCX ingestion is genuinely searchable end-to-end, retrieval
steps/citations are shown exactly when retrieval actually ran (not gated
on whether the model happens to be mocked), and the honest "nothing
indexed" fallback when no document matches.
"""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient

from app.agents.manager import build_default_agent_manager
from app.agents.task_router import TaskRouter


def _router() -> TaskRouter:
    return TaskRouter(build_default_agent_manager())


# --- Task Router: natural report/document phrasings ------------------------

_NATURAL_REPORT_QUERIES = [
    "say about the report",
    "summarize the report",
    "what is in the report?",
    "explain the report",
    "what does the report discuss?",
    "give me the main points from the report",
    "what are the findings in the report?",
]


@pytest.mark.parametrize("message", _NATURAL_REPORT_QUERIES)
def test_natural_report_queries_route_to_document_agent(message: str) -> None:
    result = _router().select(message, [])
    assert result.target == "agent"
    assert result.agent_id == "document"
    assert result.confidence >= 0.15


def test_unrelated_query_does_not_route_to_document() -> None:
    result = _router().select("what is CPU scheduling?", [])
    assert result.agent_id != "document"


# --- DOCX ingestion + grounded retrieval, end-to-end ------------------------


def _make_report_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Quarterly Maintenance Report", level=1)
    document.add_paragraph(
        "Pump station 4 underwent scheduled maintenance in March. The seal was "
        "replaced and bearings were inspected with no defects found."
    )
    document.add_paragraph(
        "Overall equipment uptime for the quarter was 98.2 percent, exceeding "
        "the 95 percent target set by operations."
    )
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


@pytest.fixture
def low_relevance_threshold(monkeypatch) -> None:
    """The app's default `rag_relevance_threshold` (0.45) is calibrated
    against real Ollama embeddings (see `app/core/config.py`) — the mock
    embedding provider these tests run against is a coarse bag-of-words
    approximation whose cosine-similarity scores for genuinely relevant
    text top out around 0.25-0.33, well under that default. Lowering the
    threshold here (not in the app's default) is what lets these tests
    exercise the *grounded* answer path against the mock provider without
    weakening the real, already-calibrated production default.
    """
    from app.core.config import get_settings

    monkeypatch.setattr(get_settings(), "rag_relevance_threshold", 0.1)


def _upload_and_ingest_report(client: TestClient) -> dict:
    uploaded = client.post(
        "/api/files/upload",
        files={
            "file": (
                "report.docx",
                io.BytesIO(_make_report_docx()),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 200, uploaded.text
    ingested = client.post("/api/documents/ingest", json={"file_id": uploaded.json()["id"]})
    assert ingested.status_code == 200, ingested.text
    return ingested.json()


def test_docx_report_is_extracted_chunked_and_indexed(client: TestClient) -> None:
    document = _upload_and_ingest_report(client)
    assert document["filename"] == "report.docx"
    assert document["doc_type"] == "docx"
    assert document["chunk_count"] >= 1

    search = client.post("/api/documents/search", json={"query": "pump station maintenance", "top_k": 3})
    assert search.status_code == 200
    results = search.json()
    assert len(results) >= 1
    assert results[0]["filename"] == "report.docx"


def test_say_about_the_report_triggers_retrieval_and_citations(client: TestClient, low_relevance_threshold) -> None:
    _upload_and_ingest_report(client)

    response = client.post("/api/chat", json={"message": "say about the report"})
    assert response.status_code == 200, response.text
    body = response.json()

    assert body["agent"]["id"] == "document"
    assert body["task_routing"]["agent_id"] == "document"

    step_ids = [s["id"] for s in body["execution"]["steps"]]
    assert "retrieving_context" in step_ids

    assert len(body["citations"]) >= 1
    assert any(c["label"].startswith("report.docx") for c in body["citations"])


def test_summarize_report_also_grounds(client: TestClient, low_relevance_threshold) -> None:
    _upload_and_ingest_report(client)

    response = client.post("/api/chat", json={"message": "what are the main points in the report?"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["agent"]["id"] == "document"
    assert len(body["citations"]) >= 1


def test_unrelated_question_uses_normal_general_behavior(client: TestClient) -> None:
    _upload_and_ingest_report(client)

    response = client.post("/api/chat", json={"message": "what is CPU scheduling?"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["agent"]["id"] != "document"
    # No knowledge-base retrieval step should appear for an unrelated request.
    step_ids = [s["id"] for s in body["execution"]["steps"]]
    assert "retrieving_context" not in step_ids


def test_nonexistent_document_query_does_not_hallucinate(client: TestClient) -> None:
    # Fresh client => empty knowledge base (isolate_data_dirs gives every
    # test its own tmp_path). No document has been ingested here at all.
    response = client.post("/api/chat", json={"message": "summarize the financial report"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["agent"]["id"] == "document"
    # No retrieval step: build_prompt short-circuits on an empty KB before
    # ever calling `.search()` — nothing to retrieve, so nothing is shown
    # as retrieved.
    step_ids = [s["id"] for s in body["execution"]["steps"]]
    assert "retrieving_context" not in step_ids
    assert body["citations"] == []
