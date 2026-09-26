"""Conversation -> document export: intent/format detection, outline
JSON parsing (retry + fallback), all four renderers, both trigger paths
(natural-language chat, and the dedicated REST endpoint), message
exclusion/empty-conversation rejection, and path confinement.

`conftest.py`'s `isolate_data_dirs` forces MODEL_PROVIDER=EMBEDDING_
PROVIDER=mock for every test here, so the outline call deterministically
falls through retry -> fixed fallback outline (the mock provider never
returns valid JSON) — this suite never depends on Ollama running on the
host.
"""

from __future__ import annotations

import pytest
from fastapi.testclient import TestClient

from app.services.conversation_export.intent import detect_export_format, is_export_request
from app.services.conversation_export.pipeline import _parse_outline_json, _slugify

# --- intent / format detection ----------------------------------------------


@pytest.mark.parametrize(
    "message,expected",
    [
        ("Can you make a ppt of this chat?", True),
        ("Please generate a document about this", True),
        ("Export this as word", True),
        ("Create an excel summary", True),
        ("Give me a pdf of our discussion", True),
        ("Please summarize this report and draft a memo.", False),
        ("Hello, how are you today?", False),
        ("There's a bug in this function.", False),
    ],
)
def test_is_export_request(message: str, expected: bool) -> None:
    assert is_export_request(message) is expected


@pytest.mark.parametrize(
    "message,expected_format",
    [
        ("make a ppt of this chat", "pptx"),
        ("generate a slideshow about this", "pptx"),
        ("export this as word", "docx"),
        ("create an excel summary", "xlsx"),
        ("give me a spreadsheet of this", "xlsx"),
        ("give me a pdf of our discussion", "pdf"),
        ("generate a document about this", "docx"),
        ("export this conversation", "docx"),  # no format noun -> default
    ],
)
def test_detect_export_format(message: str, expected_format: str) -> None:
    assert detect_export_format(message) == expected_format


# --- outline JSON parsing ----------------------------------------------------


def test_parse_outline_json_handles_fenced_json() -> None:
    raw = '```json\n{"title": "T", "subtitle": "S", "sections": [{"heading": "A", "focus": "a"}]}\n```'
    parsed = _parse_outline_json(raw)
    assert parsed is not None
    assert parsed["title"] == "T"
    assert parsed["sections"][0]["heading"] == "A"


def test_parse_outline_json_handles_bare_json_with_surrounding_prose() -> None:
    raw = 'Sure, here is the outline:\n{"title": "T", "sections": [{"heading": "A", "focus": "a"}]}\nHope that helps!'
    parsed = _parse_outline_json(raw)
    assert parsed is not None
    assert parsed["sections"][0]["heading"] == "A"


@pytest.mark.parametrize("raw", ["not json at all", "{}", '{"title": "T"}', "[]"])
def test_parse_outline_json_rejects_invalid_shapes(raw: str) -> None:
    assert _parse_outline_json(raw) is None


def test_slugify_neutralizes_path_traversal() -> None:
    assert ".." not in _slugify("../../etc/passwd")
    assert "/" not in _slugify("../../etc/passwd")


# --- outline normalization (deterministic, no LLM) ---------------------------


def test_normalize_outline_pads_two_sections_to_six_preserving_model_sections() -> None:
    from app.services.conversation_export.outline import normalize_outline

    sections, padded = normalize_outline(
        [
            {"heading": "Vibration Thresholds", "focus": "threshold values"},
            {"heading": "Inspection Schedule", "focus": "how often to inspect"},
        ]
    )
    assert len(sections) == 6
    assert padded == 4
    headings = [s["heading"] for s in sections]
    assert "Vibration Thresholds" in headings
    assert "Inspection Schedule" in headings
    # no duplicates, case/punctuation-insensitive
    assert len(headings) == len({h.lower() for h in headings})


def test_normalize_outline_empty_input_yields_six_to_eight_with_bookends() -> None:
    from app.services.conversation_export.outline import normalize_outline

    sections, padded = normalize_outline([])
    assert 6 <= len(sections) <= 8
    assert padded == len(sections)
    headings_lower = [s["heading"].lower() for s in sections]
    assert headings_lower[0] == "executive summary"
    assert "conclusion" in headings_lower[-1]


def test_normalize_outline_fallback_passthrough_stays_eight_no_padding() -> None:
    from app.services.conversation_export.outline import normalize_outline
    from app.services.conversation_export.pipeline import _DEFAULT_OUTLINE_SECTIONS

    sections, padded = normalize_outline(_DEFAULT_OUTLINE_SECTIONS)
    assert len(sections) == 8
    assert padded == 0
    assert [s["heading"] for s in sections] == [e["heading"] for e in _DEFAULT_OUTLINE_SECTIONS]


def test_normalize_outline_caps_ten_sections_keeping_bookends(caplog) -> None:
    from app.services.conversation_export.outline import normalize_outline

    raw = (
        [{"heading": "Executive Summary", "focus": "x"}]
        + [{"heading": f"Middle Topic {i}", "focus": f"topic {i}"} for i in range(8)]
        + [{"heading": "Conclusion", "focus": "wrap up"}]
    )
    assert len(raw) == 10

    sections, padded = normalize_outline(raw)
    assert len(sections) == 8
    assert padded == 0
    headings_lower = [s["heading"].lower() for s in sections]
    assert headings_lower[0] == "executive summary"
    assert headings_lower[-1] == "conclusion"


def test_normalize_outline_dedupes_case_and_punctuation_variants() -> None:
    from app.services.conversation_export.outline import normalize_outline

    sections, _ = normalize_outline(
        [
            {"heading": "Findings", "focus": "a"},
            {"heading": "findings.", "focus": "b"},
            {"heading": " FINDINGS ", "focus": "c"},
        ]
    )
    findings_like = [s for s in sections if "finding" in s["heading"].lower()]
    assert len(findings_like) == 1


def test_normalize_outline_key_results_heading_covers_findings_slot() -> None:
    from app.services.conversation_export.outline import normalize_outline

    sections, padded = normalize_outline(
        [{"heading": "Key Results", "focus": "x"}, {"heading": "Something Else", "focus": "y"}]
    )
    headings_lower = [s["heading"].lower() for s in sections]
    assert "key results" in headings_lower
    assert not any(h == "findings" for h in headings_lower)


def test_normalize_outline_orders_executive_summary_first_conclusion_last() -> None:
    from app.services.conversation_export.outline import normalize_outline

    sections, padded = normalize_outline([{"heading": "Key Discussion Points", "focus": "x"}])
    headings = [s["heading"] for s in sections]
    assert headings[0] == "Executive Summary"
    assert headings[-1] == "Conclusion"
    assert "Key Discussion Points" in headings
    assert headings.index("Key Discussion Points") < headings.index("Conclusion")


# --- export_conversation() pipeline -----------------------------------------


def _router():
    from app.models.base import ResourceClass
    from app.models.registry import ModelRegistry
    from app.models.router import ModelRouter

    return ModelRouter(ModelRegistry(), ResourceClass.SMALL)


def _sample_messages():
    from app.services.conversation_export.pipeline import ExportMessage

    return [
        ExportMessage(role="user", text="What vibration level is concerning for this turbine?"),
        ExportMessage(
            role="assistant",
            text="Above 7.1 mm/s RMS per ISO 10816 is concerning and warrants inspection.",
            citations=["iso10816_policy.pdf (p.4)"],
        ),
        ExportMessage(role="user", text="How often should we inspect it, and can you show a check script?"),
        ExportMessage(
            role="assistant",
            text="Monthly for aging units, quarterly otherwise.\n\n```python\ndef flag(v):\n    return v > 7.1\n```",
        ),
    ]


@pytest.mark.asyncio
async def test_export_conversation_falls_back_to_default_outline_under_mock(monkeypatch, tmp_path) -> None:
    from app.core.config import get_settings
    from app.services.conversation_export import export_conversation

    settings = get_settings()
    monkeypatch.setattr(settings, "generated_dir", str(tmp_path / "generated"))

    result = await export_conversation(
        _sample_messages(), "docx", model_router=_router(), settings=settings, conversation_id="conv_1"
    )

    assert result.is_mock is True
    assert result.deliverable.filename.endswith(".docx")
    generated_file = tmp_path / "generated" / result.deliverable.filename
    assert generated_file.exists()
    assert generated_file.stat().st_size == result.deliverable.size_bytes

    step_ids = [s.id for s in result.steps]
    assert "understanding_request" in step_ids
    assert "selecting_model" in step_ids
    assert "outlining_document" in step_ids
    assert any(sid.startswith("writing_section_") for sid in step_ids)
    assert "building_file" in step_ids
    assert "completed" in step_ids
    assert "no real model inference ran" in result.response_text

    outline_step = next(s for s in result.steps if s.id == "outlining_document")
    assert "(fallback)" in outline_step.detail
    section_count = sum(1 for sid in step_ids if sid.startswith("writing_section_"))
    assert 6 <= section_count <= 8


@pytest.mark.asyncio
async def test_write_section_honest_fallback_when_model_returns_nothing() -> None:
    """A padded section's LLM call may legitimately have nothing to say —
    `_write_section` must write the one honest line, never invent content
    to fill the gap.
    """
    from app.models.base import GenerationResult
    from app.services.conversation_export.pipeline import _NO_CONTENT_FALLBACK, _write_section

    class _BlankProvider:
        async def generate(self, *, model, prompt):
            return GenerationResult(text="   ", model_id=model.id, is_mock=True)

    from app.models.registry import ModelRegistry

    model = ModelRegistry().get("general-small")  # any registered ModelInfo works here
    section = await _write_section(
        _BlankProvider(), model, "Background & Context", "unrelated topic", "User: hi\n\nAssistant: hi", is_padded=True
    )
    assert section.paragraphs == [_NO_CONTENT_FALLBACK]
    assert section.bullets == []


@pytest.mark.asyncio
async def test_export_conversation_loop_budget_exhaustion_uses_honest_fallback(monkeypatch, tmp_path) -> None:
    """When the bounded loop runs out mid-way, every remaining section
    (padded or not) gets the honest one-liner instead of failing the
    export or making more LLM calls than the budget allows.
    """
    from app.core.config import get_settings
    from app.services.conversation_export import export_conversation
    from app.services.conversation_export import pipeline as pipeline_module

    settings = get_settings()
    monkeypatch.setattr(settings, "generated_dir", str(tmp_path / "generated"))
    # 1 step for the outline call + 1 step for the first section write —
    # every section after that must hit `budget.exhausted`.
    monkeypatch.setattr(pipeline_module, "_MAX_LOOP_STEPS", 2)

    result = await export_conversation(
        _sample_messages(), "docx", model_router=_router(), settings=settings, conversation_id="conv_budget"
    )

    section_steps = [s for s in result.steps if s.id.startswith("writing_section_")]
    assert len(section_steps) >= 6  # normalize_outline still guarantees 6-8 sections
    assert "budget exhausted" not in section_steps[0].detail
    assert all("budget exhausted" in s.detail for s in section_steps[1:])
    # The export must still succeed and produce a real file, never fail.
    generated_file = tmp_path / "generated" / result.deliverable.filename
    assert generated_file.exists()


@pytest.mark.asyncio
async def test_export_conversation_metadata_reflects_model_plus_padded_outline(monkeypatch, tmp_path, caplog) -> None:
    from app.core.config import get_settings
    from app.services.conversation_export import export_conversation
    from app.services.conversation_export import pipeline as pipeline_module

    settings = get_settings()
    monkeypatch.setattr(settings, "generated_dir", str(tmp_path / "generated"))

    async def fake_generate_outline(provider, model, transcript):
        return (
            {
                "title": "T",
                "subtitle": "S",
                "sections": [{"heading": "Intro", "focus": "x"}, {"heading": "Wrap Up", "focus": "y"}],
            },
            True,
            False,  # used_fallback=False -> this is a "model" outline, not the fixed fallback
        )

    monkeypatch.setattr(pipeline_module, "_generate_outline", fake_generate_outline)

    with caplog.at_level("INFO", logger="app.services.conversation_export.pipeline"):
        result = await export_conversation(
            _sample_messages(), "docx", model_router=_router(), settings=settings, conversation_id="conv_meta"
        )

    outline_step = next(s for s in result.steps if s.id == "outlining_document")
    assert "(model+padded)" in outline_step.detail

    normalized_events = [r for r in caplog.records if getattr(r, "event", None) == "outline_normalized"]
    assert normalized_events
    assert normalized_events[0].outline_source == "model+padded"
    assert normalized_events[0].outline_padded == 4


@pytest.mark.asyncio
@pytest.mark.parametrize("empty_messages", [[], None])
async def test_export_conversation_rejects_empty_conversation(monkeypatch, tmp_path, empty_messages) -> None:
    from app.core.config import get_settings
    from app.core.exceptions import InvalidRequestError
    from app.services.conversation_export import export_conversation

    settings = get_settings()
    monkeypatch.setattr(settings, "generated_dir", str(tmp_path / "generated"))

    with pytest.raises(InvalidRequestError):
        await export_conversation(
            empty_messages or [], "docx", model_router=_router(), settings=settings, conversation_id="c"
        )


@pytest.mark.asyncio
async def test_export_conversation_rejects_whitespace_only_conversation(monkeypatch, tmp_path) -> None:
    from app.core.config import get_settings
    from app.core.exceptions import InvalidRequestError
    from app.services.conversation_export import ExportMessage, export_conversation

    settings = get_settings()
    monkeypatch.setattr(settings, "generated_dir", str(tmp_path / "generated"))

    with pytest.raises(InvalidRequestError):
        await export_conversation(
            [ExportMessage(role="user", text="   "), ExportMessage(role="assistant", text="")],
            "docx",
            model_router=_router(),
            settings=settings,
            conversation_id="c",
        )


@pytest.mark.asyncio
async def test_export_conversation_extracts_code_blocks_and_references(monkeypatch, tmp_path) -> None:
    from app.core.config import get_settings
    from app.services.conversation_export.pipeline import build_references, extract_code_blocks

    settings = get_settings()
    monkeypatch.setattr(settings, "generated_dir", str(tmp_path / "generated"))

    messages = _sample_messages()
    code_blocks = extract_code_blocks(messages)
    assert len(code_blocks) == 1
    assert "def flag" in code_blocks[0].code
    assert code_blocks[0].language == "python"

    refs = build_references(messages)
    assert any(r.label == "iso10816_policy.pdf (p.4)" for r in refs)


@pytest.mark.asyncio
async def test_export_conversation_path_confinement(monkeypatch, tmp_path) -> None:
    """However `conversation_id` or model output is spelled, the written
    file must land inside `generated_path` — never escape it.
    """
    from app.core.config import get_settings
    from app.services.conversation_export import export_conversation

    settings = get_settings()
    generated_dir = tmp_path / "generated"
    monkeypatch.setattr(settings, "generated_dir", str(generated_dir))

    result = await export_conversation(
        _sample_messages(),
        "pdf",
        model_router=_router(),
        settings=settings,
        conversation_id="../../etc/passwd",
    )
    generated_file = generated_dir / result.deliverable.filename
    assert generated_file.resolve().is_relative_to(generated_dir.resolve())


# --- renderers ---------------------------------------------------------------


def _sample_spec():
    from app.services.conversation_export.spec import (
        CodeBlock,
        DocumentMetadata,
        DocumentSpec,
        QAEntry,
        ReferenceEntry,
        SectionSpec,
    )

    headings = [
        "Executive Summary",
        "Background & Context",
        "Key Discussion Points",
        "Detailed Analysis",
        "Findings / Answers Provided",
        "Decisions & Recommendations",
        "Action Items / Next Steps",
        "Conclusion",
    ]
    sections = [
        SectionSpec(heading=h, paragraphs=[f"Paragraph about {h}."], bullets=[f"Point about {h}."])
        for h in headings
    ]
    sections[0].code_blocks = [CodeBlock(language="python", code="def flag(v):\n    return v > 7.1")]

    return DocumentSpec(
        title="Turbine Maintenance Discussion",
        subtitle="Conversation summary",
        metadata=DocumentMetadata(
            conversation_id="conv_test",
            generated_at="2026-01-01T00:00:00Z",
            models_used=["mock-general"],
            message_count=4,
            generation_notice="_Structured development response — no real model inference ran._",
        ),
        sections=sections,
        qa_log=[QAEntry(question="What vibration level is concerning?", answer="Above 7.1 mm/s RMS.")],
        references=[ReferenceEntry(label="iso10816_policy.pdf (p.4)")],
    )


def test_render_docx_has_expected_headings_and_opens(tmp_path) -> None:
    import docx

    from app.services.conversation_export.renderers import render_document

    target = tmp_path / "out.docx"
    render_document(_sample_spec(), "docx", target)
    assert target.exists()

    document = docx.Document(str(target))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert "Executive Summary" in headings
    assert "Conclusion" in headings
    assert "Key Q&A Log" in headings
    assert "References" in headings


def test_render_pptx_slide_count_in_target_range(tmp_path) -> None:
    from pptx import Presentation

    from app.services.conversation_export.renderers import render_document

    target = tmp_path / "out.pptx"
    render_document(_sample_spec(), "pptx", target)
    assert target.exists()

    prs = Presentation(str(target))
    slide_count = len(prs.slides._sldIdLst)
    assert 8 <= slide_count <= 10


def test_render_xlsx_has_expected_sheets(tmp_path) -> None:
    from openpyxl import load_workbook

    from app.services.conversation_export.renderers import render_document

    target = tmp_path / "out.xlsx"
    render_document(_sample_spec(), "xlsx", target)
    assert target.exists()

    workbook = load_workbook(str(target))
    assert workbook.sheetnames == ["Summary", "Key Points", "Q&A Log", "References"]


def test_render_pdf_has_multiple_pages(tmp_path) -> None:
    from pypdf import PdfReader

    from app.services.conversation_export.renderers import render_document

    target = tmp_path / "out.pdf"
    render_document(_sample_spec(), "pdf", target)
    assert target.exists()

    reader = PdfReader(str(target))
    assert len(reader.pages) > 1


# --- API: dedicated REST endpoint --------------------------------------------


def test_export_endpoint_docx_roundtrip(client: TestClient) -> None:
    response = client.post(
        "/api/conversations/conv_abc/export",
        json={
            "format": "docx",
            "messages": [
                {"role": "user", "text": "What vibration level is concerning?"},
                {"role": "assistant", "text": "Above 7.1 mm/s RMS is concerning."},
            ],
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["deliverable"]["filename"].endswith(".docx")
    assert body["deliverable"]["status"] == "ready"
    assert body["execution"]["steps"]

    download = client.get(f"/api/deliverables/{body['deliverable']['id']}")
    assert download.status_code == 200
    assert download.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml"
    )


def test_export_endpoint_rejects_empty_conversation(client: TestClient) -> None:
    response = client.post("/api/conversations/conv_empty/export", json={"format": "docx", "messages": []})
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "invalid_request"


def test_export_endpoint_rejects_bad_format(client: TestClient) -> None:
    response = client.post(
        "/api/conversations/conv_x/export",
        json={"format": "exe", "messages": [{"role": "user", "text": "hi"}]},
    )
    assert response.status_code == 400


@pytest.mark.parametrize("fmt", ["pptx", "xlsx", "pdf"])
def test_export_endpoint_supports_every_format(client: TestClient, fmt: str) -> None:
    response = client.post(
        "/api/conversations/conv_multi/export",
        json={
            "format": fmt,
            "messages": [
                {"role": "user", "text": "Summarize our maintenance discussion."},
                {"role": "assistant", "text": "Sure — here is a summary of what we covered."},
            ],
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["deliverable"]["filename"].endswith(f".{fmt}")


# --- API: natural-language chat trigger --------------------------------------


def test_chat_nl_trigger_routes_to_export_with_history(client: TestClient) -> None:
    response = client.post(
        "/api/chat",
        json={
            "message": "Can you generate a document about this conversation?",
            "conversation_history": [
                {"role": "user", "text": "What vibration level is concerning for this turbine?"},
                {"role": "assistant", "text": "Above 7.1 mm/s RMS per ISO 10816 is concerning."},
            ],
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["agent"]["id"] == "document"
    assert body["task_routing"]["reason"] == "The request asks to export this conversation as a document."
    assert len(body["deliverables"]) == 1
    assert body["deliverables"][0]["filename"].endswith(".docx")


def test_chat_nl_trigger_excludes_the_trigger_message_itself(client: TestClient, monkeypatch) -> None:
    """The export request ("make a ppt of this") must never appear inside
    the generated document's transcript — only prior history does.
    """
    captured: dict[str, list[str]] = {}
    from app.services.conversation_export import pipeline as pipeline_module

    original = pipeline_module.build_transcript_text

    def spy(messages):
        captured["texts"] = [m.text for m in messages]
        return original(messages)

    monkeypatch.setattr(pipeline_module, "build_transcript_text", spy)

    response = client.post(
        "/api/chat",
        json={
            "message": "make a ppt of this chat",
            "conversation_history": [
                {"role": "user", "text": "What vibration level is concerning?"},
                {"role": "assistant", "text": "Above 7.1 mm/s RMS is concerning."},
            ],
        },
    )
    assert response.status_code == 200, response.text
    assert "make a ppt of this chat" not in captured["texts"]
    assert any("vibration" in t for t in captured["texts"])


def test_chat_nl_trigger_requires_history(client: TestClient) -> None:
    response = client.post("/api/chat", json={"message": "generate a document about this"})
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "invalid_request"


def test_chat_manual_mode_never_triggers_export(client: TestClient) -> None:
    """A concrete agent_id must skip the Task Router (and therefore export
    intent detection) entirely — unchanged Segment 2/3 behavior.
    """
    response = client.post(
        "/api/chat",
        json={
            "agent_id": "document",
            "message": "make a ppt of this chat",
            "conversation_history": [{"role": "user", "text": "hello"}],
        },
    )
    assert response.status_code == 200
    assert response.json()["task_routing"] is None
    assert response.json()["deliverables"] == []
